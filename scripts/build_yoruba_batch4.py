import csv
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Yoruba_targeted_batch3_RECORDING_SHEET.docx"
DATA = ROOT / "work" / "yoruba_batch4_items.csv"
DOCX_OUT = ROOT / "outputs" / "Yoruba_targeted_batch4_RECORDING_SHEET.docx"
CSV_OUT = ROOT / "outputs" / "yoruba_targeted_batch4_manifest.csv"

with DATA.open("r", encoding="utf-8-sig", newline="") as fh:
    items = list(csv.DictReader(fh))

if len(items) != 30:
    raise ValueError(f"Expected 30 items, found {len(items)}")

expected = [f"speaker01_{number:04d}.m4a" for number in range(191, 221)]
actual = [item["filename"] for item in items]
if actual != expected:
    raise ValueError("Batch 4 filenames are not the exact 0191-0220 sequence")

CSV_OUT.write_text(DATA.read_text(encoding="utf-8-sig"), encoding="utf-8-sig")

doc = Document(SOURCE)
doc.paragraphs[0].runs[0].text = "Yoruba Voice Research - Batch 4"
doc.paragraphs[1].runs[0].text = "Voice-assistant commands, meaning contrasts, emphasis, emotion, and connected speech"

table = doc.tables[0]
if len(table.rows) != 31:
    raise ValueError("The retained recording-sheet template does not contain 30 item rows")

for row, item in zip(table.rows[1:], items):
    values = (item["filename"], item["yoruba"], item["english_meaning"])
    for cell, value in zip(row.cells, values):
        paragraph = cell.paragraphs[0]
        if not paragraph.runs:
            paragraph.add_run()
        paragraph.runs[0].text = value
        for extra in paragraph.runs[1:]:
            extra._element.getparent().remove(extra._element)

doc.sections[0].header.paragraphs[0].runs[0].text = "Yoruba Voice Research | Batch 4"
doc.sections[0].footer.paragraphs[0].runs[0].text = "speaker01_0191-0220 | 30 sentences"
doc.core_properties.title = "Yoruba Voice Research - Batch 4 Recording Sheet"
doc.core_properties.subject = "Controlled Yoruba voice-model evaluation recordings"
doc.save(DOCX_OUT)

print(DOCX_OUT)
print(CSV_OUT)
