from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK_OUT = ROOT / "work" / "Yoruba_recording_sheet.docx"
USER_OUT = ROOT / "outputs" / "Yoruba_recording_sheet.docx"

PROMPTS = [
    ("speaker01_0001", "Ẹ káàárọ̀.", "Good morning."),
    ("speaker01_0002", "Báwo ni o ṣe wà?", "How are you?"),
    ("speaker01_0003", "Mo wà dáadáa.", "I am fine."),
    ("speaker01_0004", "Mo fẹ́ lọ sí ilé.", "I want to go home."),
    ("speaker01_0005", "Jọ̀wọ́, fún mi ní omi.", "Please give me water."),
    ("speaker01_0006", "Ọmọ náà ń ka ìwé.", "That child is reading a book."),
    ("speaker01_0007", "Ṣé o ti jẹun?", "Have you eaten?"),
    ("speaker01_0008", "A ó pàdé lọ́la.", "We will meet tomorrow."),
    ("speaker01_0009", "Ọjọ́ náà dára gan-an.", "The day was very good."),
    ("speaker01_0010", "Ẹ jọ̀wọ́, ẹ sọ̀rọ̀ díẹ̀díẹ̀.", "Please speak slowly."),
    ("speaker01_0011", "Ọ̀rẹ́ mi fẹ́ ra ẹja ní ọjà.", "My friend wants to buy fish at the market."),
    ("speaker01_0012", "Ṣadé ń ṣe oúnjẹ ní ilé.", "Sade is cooking at home."),
    ("speaker01_0013", "Àwọn ọmọ ń lọ sí ilé ẹ̀kọ́.", "The children are going to school."),
    ("speaker01_0014", "Òjò ń rọ̀ lónìí.", "It is raining today."),
    ("speaker01_0015", "Mo gbọ́ ohun tí o sọ.", "I heard what you said."),
    ("speaker01_0016", "Ṣé o lè ràn mí lọ́wọ́?", "Can you help me?"),
    ("speaker01_0017", "Jọ̀wọ́, pa ilẹ̀kùn náà.", "Please close that door."),
    ("speaker01_0018", "Ẹ̀gbọ́n mi ń ṣiṣẹ́ ní ọ́fíìsì.", "My older sibling works in an office."),
    ("speaker01_0019", "Mo ní ìpàdé lọ́la.", "I have a meeting tomorrow."),
    ("speaker01_0020", "Inú mi dùn láti rí ẹ!", "I am happy to see you!"),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.right_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 24, "0B2545", 0, 8),
    ("Heading 1", 16, "2E74B5", 18, 10),
    ("Heading 2", 13, "2E74B5", 14, 7),
    ("Heading 3", 12, "1F4D78", 10, 5),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph(style="Title")
title.add_run("Yoruba Voice Calibration").bold = True
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
r = subtitle.add_run("20-sentence computer recording sheet")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string("556070")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Purpose: ")
r.bold = True
p.add_run("Make a short test recording before the full research dataset. Check and correct the Yoruba text before recording.")

doc.add_heading("Record with your Windows computer", level=1)
steps = [
    ("Step 1 - Prepare", "Go to a quiet room. Turn off fans, music, television, and computer notifications."),
    ("Step 2 - Open the recorder", "Open the Windows Start menu, search for Sound Recorder, and open it. Voice Recorder is the older name on some computers."),
    ("Step 3 - Select the microphone", "If needed, open Windows Settings > System > Sound > Input and choose the microphone you want to use. Keep your mouth about 15-20 cm from it."),
    ("Step 4 - Make one file per sentence", "Press Record, wait one second, read one sentence naturally, wait one second, and stop. Do not exaggerate the tones."),
    ("Step 5 - Rename the recording", "Use the filename shown beside the sentence. Your computer may save M4A instead of WAV; that is acceptable for this first calibration."),
    ("Step 6 - Send the recordings", "Put all 20 recordings in one folder, compress the folder to a ZIP file, and attach the ZIP to the Codex task."),
]
for label, body in steps:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    rr = p.add_run(label + ": ")
    rr.bold = True
    p.add_run(body)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(8)
note.paragraph_format.space_after = Pt(12)
rr = note.add_run("Important: ")
rr.bold = True
rr.font.color.rgb = RGBColor.from_string("7A5A00")
note.add_run("If a Yoruba sentence or tone mark is wrong or unnatural, correct it in this document before reading it. Your correction becomes the authoritative transcript.")

doc.add_heading("Sentences to verify and record", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Filename"
hdr[1].text = "Verify, then read this sentence"
for cell in hdr:
    shade(cell, "E8EEF5")
    for run in cell.paragraphs[0].runs:
        run.bold = True
for stem, yoruba, gloss in PROMPTS:
    cells = table.add_row().cells
    cells[0].text = stem
    p = cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    yr = p.add_run(yoruba)
    yr.bold = True
    yr.font.size = Pt(12)
    gloss_p = cells[1].add_paragraph(gloss)
    gloss_p.paragraph_format.space_after = Pt(0)
    gloss_run = gloss_p.runs[0]
    gloss_run.italic = True
    gloss_run.font.size = Pt(9)
    gloss_run.font.color.rgb = RGBColor.from_string("556070")
set_table_geometry(table, [2300, 7060])
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

doc.add_heading("When you finish", level=1)
p = doc.add_paragraph()
p.add_run("Attach one ZIP containing the 20 audio files. ").bold = True
p.add_run("Codex will check sound quality, clipping, background noise, file matching, and whether the setup is suitable for the full study.")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Yoruba voice research - calibration")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("777777")

WORK_OUT.parent.mkdir(parents=True, exist_ok=True)
USER_OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(WORK_OUT)
USER_OUT.write_bytes(WORK_OUT.read_bytes())
print(WORK_OUT)
print(USER_OUT)
