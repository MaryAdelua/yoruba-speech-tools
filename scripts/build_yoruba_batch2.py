import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
DOCX_OUT = OUT / "Yoruba_targeted_batch2_RECORDING_SHEET.docx"
CSV_OUT = OUT / "yoruba_targeted_batch2_manifest.csv"

ITEMS = [
    ("speaker01_0131.m4a", "Ará Èkó ni Bọ́lá.", "Bola is a resident of Lagos.", "lexical-tone contrast: ará"),
    ("speaker01_0132.m4a", "Ara ọmọ náà gbóná.", "The child's body is hot.", "lexical-tone contrast: ara"),
    ("speaker01_0133.m4a", "Àrá ń sán lókè ọ̀run.", "Thunder is rumbling in the sky.", "lexical-tone contrast: àrá"),
    ("speaker01_0134.m4a", "Ogun náà parí ní àná.", "The war ended yesterday.", "tone sequence: ogun"),
    ("speaker01_0135.m4a", "Dókítà fún mi ní oògùn.", "The doctor gave me medicine.", "tone sequence: oògùn"),
    ("speaker01_0136.m4a", "Kúnlé jogún ilé bàbá rẹ̀.", "Kunle inherited his father's house.", "tone sequence: jogún"),
    ("speaker01_0137.m4a", "Bàbá mi ń ṣe òwò ní ọjà.", "My father trades at the market.", "lexical-tone contrast: òwò"),
    ("speaker01_0138.m4a", "Owó náà wà nínú àpò mi.", "The money is in my bag.", "lexical-tone contrast: owó"),
    ("speaker01_0139.m4a", "Ọwọ́ mi mọ́.", "My hand is clean.", "lexical-tone contrast: ọwọ́"),
    ("speaker01_0140.m4a", "Mo dé ní ìgbà tó yẹ.", "I arrived at the right time.", "lexical-tone contrast: ìgbà"),
    ("speaker01_0141.m4a", "Ìyá gbé igbá náà kalẹ̀.", "Mother put the calabash down.", "lexical-tone contrast: igbá"),
    ("speaker01_0142.m4a", "Mo kọ́ ẹ̀kọ́ tuntun lónìí.", "I learned a new lesson today.", "lexical-tone contrast: ẹ̀kọ́"),
    ("speaker01_0143.m4a", "Ọmọ náà jẹ ẹ̀kọ ní òwúrọ̀.", "The child ate corn pap in the morning.", "lexical-tone contrast: ẹ̀kọ"),
    ("speaker01_0144.m4a", "Bọ́lá lọ sí ilé ẹ̀kọ́.", "Bola went to school.", "compound tone sequence"),
    ("speaker01_0145.m4a", "Kúnlé lọ sí ọjà.", "Kunle went to the market.", "affirmative statement"),
    ("speaker01_0146.m4a", "Kúnlé ò lọ sí ọjà.", "Kunle did not go to the market.", "negation contrast"),
    ("speaker01_0147.m4a", "Ó ti dé ilé.", "He or she has arrived home.", "completed-action statement"),
    ("speaker01_0148.m4a", "Kò tíì dé ilé.", "He or she has not arrived home yet.", "negative aspect contrast"),
    ("speaker01_0149.m4a", "Ṣé Bọ́lá ti jí?", "Has Bola woken up?", "yes-no question intonation"),
    ("speaker01_0150.m4a", "Bọ́lá ti jí.", "Bola has woken up.", "question-statement contrast"),
    ("speaker01_0151.m4a", "Adé ni ó ra ìwé náà.", "It was Ade who bought the book.", "focus construction"),
    ("speaker01_0152.m4a", "Adé ra ìwé náà.", "Ade bought the book.", "neutral-focus contrast"),
    ("speaker01_0153.m4a", "Ẹ má ṣe bẹ̀rù.", "Do not be afraid.", "negative command"),
    ("speaker01_0154.m4a", "Ẹ jọ̀wọ́, ẹ tún un sọ.", "Please say it again.", "polite request and connected speech"),
    ("speaker01_0155.m4a", "Mo fẹ́ kí o wá ní báyìí.", "I want you to come now.", "connected speech"),
    ("speaker01_0156.m4a", "Nígbà tí mo dé, wọ́n ti lọ.", "By the time I arrived, they had left.", "tone in a complex sentence"),
    ("speaker01_0157.m4a", "Bí òjò bá rọ̀, a ó dúró sí ilé.", "If it rains, we will stay at home.", "conditional and connected speech"),
    ("speaker01_0158.m4a", "Ìpàdé náà bẹ̀rẹ̀ ní aago mẹ́wàá.", "The meeting began at ten o'clock.", "meeting term and numeral"),
    ("speaker01_0159.m4a", "Mo ra ọ̀gẹ̀dẹ̀ méje ní ọjà.", "I bought seven bananas at the market.", "numeral and vowel coverage"),
    ("speaker01_0160.m4a", "Inú mi kò dùn sí ohun tí ó ṣẹlẹ̀.", "I am unhappy about what happened.", "expressive negative speech"),
]


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.first_child_found_in("w:tcMar")
    if old is not None:
        tc_pr.remove(old)
    tc_mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    old_ind = tbl_pr.first_child_found_in("w:tblInd")
    if old_ind is not None:
        tbl_pr.remove(old_ind)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_font(run, size, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


OUT.mkdir(parents=True, exist_ok=True)
with CSV_OUT.open("w", newline="", encoding="utf-8-sig") as fh:
    writer = csv.writer(fh)
    writer.writerow(["filename", "yoruba", "english_meaning", "research_purpose"])
    writer.writerows(ITEMS)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

# compact_reference_guide token map
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(8)
run = title.add_run("Yoruba Voice Research - Batch 2")
set_font(run, 23, bold=True, color="0B2545")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
run = subtitle.add_run("Tone contrasts, meaning preservation, grammar, and connected speech")
set_font(run, 12, color="4F6272")

lead = doc.add_paragraph()
lead.paragraph_format.space_after = Pt(8)
r = lead.add_run("Review first: ")
set_font(r, 11, bold=True)
r = lead.add_run("Read the Yoruba and English silently. If any sentence does not sound natural to you, do not record it yet; tell Codex the filename and your correction.")
set_font(r, 11)

instructions = [
    ("Record", "one Yoruba sentence per file; do not read the filename or English."),
    ("Timing", "wait about half a second before and after speaking."),
    ("Delivery", "rename each recording exactly as shown, then ZIP all 30 files."),
]
for label, detail in instructions:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.187)
    p.paragraph_format.first_line_indent = Inches(-0.187)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ": ")
    set_font(r, 10.5, bold=True)
    r = p.add_run(detail)
    set_font(r, 10.5)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = ["Exact filename", "Read this Yoruba sentence", "English meaning - do not read"]
for cell, label in zip(table.rows[0].cells, headers):
    cell.text = label
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "E8EEF5")
    cell._tc.get_or_add_tcPr().append(shd)
    set_font(cell.paragraphs[0].runs[0], 9.5, bold=True, color="0B2545")
repeat = OxmlElement("w:tblHeader")
repeat.set(qn("w:val"), "true")
table.rows[0]._tr.get_or_add_trPr().append(repeat)

for filename, yoruba, english, purpose in ITEMS:
    cells = table.add_row().cells
    cells[0].text = filename
    cells[1].text = yoruba
    cells[2].text = english
    set_font(cells[0].paragraphs[0].runs[0], 8.5)
    set_font(cells[1].paragraphs[0].runs[0], 11.5, bold=True)
    set_font(cells[2].paragraphs[0].runs[0], 9)
    for cell in cells:
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.08

set_table_geometry(table, [1900, 4200, 3260])

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("Yoruba Voice Research | Batch 2")
set_font(r, 8, color="777777")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("speaker01_0131-0160 | 30 sentences")
set_font(r, 8, color="777777")

doc.core_properties.title = "Yoruba Voice Research - Batch 2 Recording Sheet"
doc.core_properties.subject = "Controlled Yoruba voice-model evaluation recordings"
doc.core_properties.author = "Yoruba Voice Research Project"
doc.save(DOCX_OUT)

print(DOCX_OUT)
print(CSV_OUT)
